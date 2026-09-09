"""Le a aba "catalog" de uma planilha de mapeamento de catalogo (Ericsson/Huawei/Nokia) e
converte cada linha (par tabela/coluna) para o formato descrito em `modelo.json`, na mesma
pasta deste arquivo.

Standalone por decisao explicita: nao importa nada do resto do projeto
`mapeamento_planilha_3_4` (nem `models.py`, nem `loaders/`) -- so depende de `openpyxl` e da
biblioteca padrao do Python, para poder ser copiado e usado por outra pessoa sem precisar do
restante do repositorio.

Uso:
    python parse_mapping_catalog.py <planilha.xlsx>
    python parse_mapping_catalog.py <pasta_com_varias_planilhas>

Em ambos os casos, para cada planilha .xlsx encontrada, escreve um arquivo
"<nome_da_planilha>__analise_felipe.json" ao lado do arquivo original, com o resultado de
`parse_catalog_xlsx`.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import TypedDict

from openpyxl import load_workbook


class ColumnMapping(TypedDict):
    bdcolname: str
    managedObjectFullPathField: bool
    managedObjectNameField: bool
    entityField: bool
    udn3: str


class TableMapping(TypedDict):
    VERSION: str
    dbn0Insert: bool
    objectType: str
    entityFields: list[str]
    columns: dict[str, ColumnMapping]


CatalogMapping = dict[str, TableMapping]


def parse_catalog_xlsx(xlsx_path: Path) -> CatalogMapping:
    """Le a aba `catalog` de `xlsx_path` e retorna o mapeamento tableName -> TableMapping.

    Cada linha da planilha e um par tabela/coluna. Linhas sem `id` (sobra do range usado na
    planilha) sao ignoradas. Linhas com `id` mas sem `bdcolname` representam uma tabela sem
    nenhuma coluna mapeada -- ainda geram a entrada da tabela (com `columns` vazio), so nao
    geram entrada em `columns`.

    `managedObjectNameField` normalmente e um flag `"X"`/`"-"`, mas a planilha Ericsson repete
    o `id` da tabela nessa coluna em toda linha, em vez do flag -- nesse caso o valor vira
    `True` so na linha cujo `bdcolname` bate com o valor repetido, `False` nas demais.
    """
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        worksheet = workbook["catalog"]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows)
        catalog: CatalogMapping = {}
        for row in rows:
            data = dict(zip(header, row))
            table_id = data.get("id")
            if table_id is None:
                continue

            table_name = data.get("tableName") or table_id
            table = catalog.get(table_name)
            if table is None:
                table = {
                    "VERSION": _as_str(data.get("VERSION")),
                    "dbn0Insert": _as_bool_insert(data.get("dbn0Insert")),
                    "objectType": _as_str(data.get("objectType")),
                    "entityFields": [],
                    "columns": {},
                }
                catalog[table_name] = table

            bdcolname = data.get("bdcolname")
            if bdcolname is None:
                continue

            is_entity_field = data.get("entityField") == "X"
            if is_entity_field:
                table["entityFields"].append(bdcolname)

            table["columns"][bdcolname] = {
                "bdcolname": bdcolname,
                "managedObjectFullPathField": data.get("managedObjectFullPathField") == "X",
                "managedObjectNameField": _is_managed_object_name_field(data, bdcolname),
                "entityField": is_entity_field,
                "udn3": _as_str(data.get("udn3")),
            }
        return catalog
    finally:
        workbook.close()


def _is_managed_object_name_field(data: dict[str, object], bdcolname: str) -> bool:
    value = data.get("managedObjectNameField")
    if value in ("X", "-", None):
        return value == "X"
    return value == bdcolname


def _as_str(value: object) -> str:
    return "" if value is None else str(value)


def _as_bool_insert(value: object) -> bool:
    return str(value).strip().lower() == "true"


def iter_catalog_xlsx_files(path: Path) -> list[Path]:
    """Se `path` for um arquivo, retorna so ele; se for uma pasta, busca .xlsx
    recursivamente (ignora arquivos de lock do Excel, que comecam com `~$`)."""
    if path.is_file():
        return [path]
    return sorted(p for p in path.rglob("*.xlsx") if not p.name.startswith("~$"))


def write_analysis_json(xlsx_path: Path) -> Path:
    """Roda `parse_catalog_xlsx` em `xlsx_path` e escreve o resultado ao lado do arquivo, com
    o mesmo nome + sufixo `__analise_felipe.json`. Retorna o caminho escrito."""
    catalog = parse_catalog_xlsx(xlsx_path)
    output_path = xlsx_path.with_name(f"{xlsx_path.stem}__analise_felipe.json")
    output_path.write_text(json.dumps(catalog, indent=2, ensure_ascii=False), encoding="utf-8")
    return output_path


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Le uma planilha de mapeamento de catalogo (ou uma pasta com varias) e gera, ao "
            "lado de cada .xlsx, um <nome>__analise_felipe.json no formato de modelo.json."
        )
    )
    parser.add_argument("path", type=Path, help="Planilha .xlsx ou pasta com varias planilhas")
    args = parser.parse_args(argv)

    xlsx_files = iter_catalog_xlsx_files(args.path)
    if not xlsx_files:
        print(f"Nenhum .xlsx encontrado em {args.path}", file=sys.stderr)
        raise SystemExit(1)

    for xlsx_path in xlsx_files:
        output_path = write_analysis_json(xlsx_path)
        print(f"{xlsx_path.name} -> {output_path.name}")


if __name__ == "__main__":
    main()
